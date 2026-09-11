# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/bonzinilabs/docs/cargo/outils/Bonzini-Comparateur-Cotations-Fret.xlsx"

VIOLET = "6B2FFA"; AMBER = "FFA31A"; ORANGE = "FF4D1A"
GREY_BG = "F2F2F6"; VIOLET_BG = "EFE9FF"; AMBER_BG = "FFF6E6"
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")
BLUE = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
H1 = Font(name="Arial", size=16, bold=True, color=VIOLET)
H2 = Font(name="Arial", size=11, bold=True, color="FFFFFF")
SMALL = Font(name="Arial", size=8, color="6B6B78")
thin = Side(style="thin", color="DCDCE4")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ═══════════ ONGLET 1 — MODE D'EMPLOI ═══════════
ws = wb.active; ws.title = "Mode d'emploi"
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 105
ws['B2'] = "COMPARATEUR DE COTATIONS — FRET MARITIME"; ws['B2'].font = H1
ws['B3'] = "Bonzini Labs · Division Cargo"; ws['B3'].font = Font(name="Arial", size=10, color=AMBER)

rows = [
 ("", ""),
 ("À QUOI SERT CE FICHIER", "h"),
 ("Comparer plusieurs offres de transport maritime pour un même conteneur, sur une base honnête.", "t"),
 ("Le piège du métier : une offre annoncée moins chère peut coûter plus cher une fois toutes les", "t"),
 ("surcharges ajoutées. Ce fichier force la comparaison sur le TOTAL, pas sur le prix affiché.", "t"),
 ("", ""),
 ("LES 4 ONGLETS", "h"),
 ("1. Mode d'emploi — cette page", "t"),
 ("2. Avant de commencer — ce qu'il faut préparer avant d'ouvrir les sites", "t"),
 ("3. Comparateur — la grille à remplir avec les offres reçues", "t"),
 ("4. Historique — le journal de toutes tes cotations, conteneur après conteneur", "t"),
 ("", ""),
 ("CODE COULEUR", "h"),
 ("Fond JAUNE  = cellule à remplir par toi", "t"),
 ("Texte BLEU  = valeur que tu saisis", "t"),
 ("Texte NOIR  = calcul automatique, ne pas modifier", "t"),
 ("", ""),
 ("COMMENT L'UTILISER", "h"),
 ("Étape 1 — Ouvre l'onglet « Avant de commencer » et rassemble tout ce qui est listé.", "t"),
 ("Étape 2 — Crée tes comptes chez les armateurs (liens dans l'onglet 2).", "t"),
 ("Étape 3 — Demande la même cotation à chacun, avec EXACTEMENT les mêmes paramètres.", "t"),
 ("Étape 4 — Reporte chaque offre dans l'onglet « Comparateur », ligne par ligne.", "t"),
 ("Étape 5 — Compare les TOTAUX, pas les prix de base. Puis regarde la section C.", "t"),
 ("Étape 6 — Reporte le résultat dans l'onglet « Historique ». C'est ta base de données prix.", "t"),
 ("", ""),
 ("LA RÈGLE D'OR", "h"),
 ("On ne compare jamais deux prix. On compare deux TOTAUX, rendus au même endroit.", "t"),
 ("", ""),
 ("QUESTION À POSER SYSTÉMATIQUEMENT", "h"),
 ("« Ce prix est-il tout compris jusqu'au terminal de Douala ? Si non, donnez-moi la liste", "t"),
 ("complète des surcharges, et dites-moi jusqu'à quelle date le prix est garanti. »", "t"),
]
r = 5
for txt, kind in rows:
    if kind == "h":
        ws.cell(r, 2, txt).font = Font(name="Arial", size=10, bold=True, color=VIOLET)
        ws.cell(r, 2).fill = PatternFill("solid", fgColor=VIOLET_BG)
    elif kind == "t":
        ws.cell(r, 2, txt).font = BLACK
    r += 1

# ═══════════ ONGLET 2 — AVANT DE COMMENCER ═══════════
ws = wb.create_sheet("Avant de commencer")
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 42
ws.column_dimensions['C'].width = 62
ws['B2'] = "AVANT DE COMMENCER"; ws['B2'].font = H1
ws['B3'] = "Rassemble tout ceci AVANT d'ouvrir les sites — sinon tu devras t'interrompre."; ws['B3'].font = SMALL

def section(ws, r, title, color=VIOLET):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = ws.cell(r, 2, title); c.font = H2
    c.fill = PatternFill("solid", fgColor=color); c.alignment = Alignment(vertical="center")
    ws.row_dimensions[r].height = 20
    return r + 1

r = 5
r = section(ws, r, "  1 · DOCUMENTS POUR CRÉER LE COMPTE (vérification KYC)")
kyc = [
 ("Lettre d'approbation", "Sur papier à en-tête de l'entreprise, autorisant la personne à ouvrir le compte"),
 ("Pièce d'identité entreprise", "Copie de la carte d'identité d'entreprise ou carte de visite"),
 ("Justificatif d'adresse", "Facture d'électricité ou de téléphone au nom ET à l'adresse de l'entreprise"),
 ("Preuve d'existence légale", "Licence commerciale, registre de commerce, attestation TVA, certificat fiscal, ou licence import-export"),
 ("Adresse complète", "L'adresse postale complète de l'entreprise"),
 ("Numéro fiscal", "Le numéro d'identification fiscale (NIU au Cameroun)"),
 ("Email professionnel", "Le lien de validation expire en 48h — prévois d'être disponible"),
]
for a, b in kyc:
    ws.cell(r, 2, a).font = BOLD; ws.cell(r, 2).border = BOX
    ws.cell(r, 3, b).font = BLACK; ws.cell(r, 3).border = BOX
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")
    r += 1

r += 1
r = section(ws, r, "  2 · DÉCISION À PRENDRE AVANT — quelle entité ouvre le compte ?", ORANGE)
ws.cell(r, 2, "Entité chinoise").font = BOLD
ws.cell(r, 3, "C'est elle qui charge le conteneur à Guangzhou. C'est donc le chargeur naturel sur le connaissement.").font = BLACK
ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top"); r += 1
ws.cell(r, 2, "Norton Goss Bonzini SARL").font = BOLD
ws.cell(r, 3, "Entité camerounaise. C'est elle qui reçoit à Douala — le destinataire.").font = BLACK
ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top"); r += 1
ws.cell(r, 2, "À TRANCHER").font = Font(name="Arial", size=10, bold=True, color=ORANGE)
ws.cell(r, 3, "Ouvre d'abord un compte avec l'entité qui a les documents les plus complets. Tu pourras ajouter l'autre ensuite. Ce choix sera repris au module 13.").font = BLACK
ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top"); r += 2

r = section(ws, r, "  3 · LES DONNÉES DE TA SIMULATION (dossier BZ-CT-2604-01)", AMBER)
fiche = [
 ("Port de départ", "Nansha (Guangzhou), Chine", 0),
 ("Port d'arrivée", "Douala, Cameroun", 0),
 ("Type de conteneur", "40' High Cube (40HC)", 0),
 ("Nombre de conteneurs", "1", 0),
 ("Date de départ souhaitée", "à saisir", 1),
 ("Description marchandise", "Marchandise générale : papier, adhésifs, textile, quincaillerie, pièces détachées, cosmétiques", 0),
 ("Poids brut total (kg)", "17 710", 1),
 ("Volume total (m3)", "64", 1),
 ("MARCHANDISE DANGEREUSE ?", "OUI — UN3481, Classe 9 (batteries lithium dans les téléphones)", 0),
 ("Type de service souhaité", "CY/CY (conteneur sorti plein du port de Douala)", 0),
 ("Incoterm", "à définir avec le client", 1),
 ("Nom du chargeur (shipper)", "à saisir", 1),
 ("Nom du destinataire (consignee)", "à saisir", 1),
]
for a, b, editable in fiche:
    ws.cell(r, 2, a).font = BOLD; ws.cell(r, 2).border = BOX
    c = ws.cell(r, 3, b); c.border = BOX
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if editable: c.font = BLUE; c.fill = INPUT_FILL
    else: c.font = BLACK
    r += 1

r += 1
r = section(ws, r, "  4 · LE TEST MALIN — fais DEUX simulations", ORANGE)
ws.cell(r, 2, "Simulation A").font = BOLD
ws.cell(r, 3, "SANS déclarer la marchandise dangereuse (marchandise générale seule)").font = BLACK; r += 1
ws.cell(r, 2, "Simulation B").font = BOLD
ws.cell(r, 3, "AVEC la marchandise dangereuse déclarée (UN3481, classe 9)").font = BLACK; r += 1
ws.cell(r, 2, "L'écart =").font = Font(name="Arial", size=10, bold=True, color=ORANGE)
ws.cell(r, 3, "Le coût réel de tes téléphones dans ce conteneur. Personne ne te donnera ce chiffre autrement.").font = BLACK
ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top"); r += 2

r = section(ws, r, "  5 · OÙ ALLER")
sites = [
 ("Maersk — inscription", "accounts.maersk.com/ocean-maeu/auth/register"),
 ("Maersk — produit Spot", "maersk.com/transportation-services/maersk-spot"),
 ("CMA CGM — SpotOn", "cma-cgm.fr/my-cma-cgm/prices/cotation-spoton"),
 ("MSC", "msc.com"),
 ("COSCO", "lines.coscoshipping.com"),
]
for a, b in sites:
    ws.cell(r, 2, a).font = BOLD; ws.cell(r, 2).border = BOX
    ws.cell(r, 3, b).font = Font(name="Arial", size=10, color=VIOLET); ws.cell(r, 3).border = BOX
    r += 1

r += 1
r = section(ws, r, "  6 · ORDRE DE GRANDEUR — pour vérifier que ta cotation est plausible", ORANGE)
ws.cell(r, 2, "Conteneur 20' Chine->Douala").font = BOLD
ws.cell(r, 3, "environ 1 800 à 4 500 USD").font = BLACK; r += 1
ws.cell(r, 2, "Conteneur 40' Chine->Douala").font = BOLD
ws.cell(r, 3, "environ 2 200 à 5 500 USD").font = BLACK; r += 1
ws.cell(r, 2, "Durée de transit").font = BOLD
ws.cell(r, 3, "environ 25 à 45 jours").font = BLACK; r += 1
ws.cell(r, 2, "ATTENTION").font = Font(name="Arial", size=10, bold=True, color=ORANGE)
ws.cell(r, 3, "Fourchettes de marché relevées en 2026, NON OFFICIELLES et très volatiles. Elles servent uniquement à détecter une cotation aberrante — jamais à construire un prix client.").font = BLACK
ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical="top")

# ═══════════ ONGLET 3 — COMPARATEUR ═══════════
ws = wb.create_sheet("Comparateur")
ws.column_dimensions['A'].width = 44
for col in "BCDEF": ws.column_dimensions[col].width = 17
ws['A1'] = "COMPARATEUR DE COTATIONS"; ws['A1'].font = H1

hdr = [("A3","Référence dossier","B3","BZ-CT-2604-01"),
       ("A4","Date de la cotation","B4","JJ/MM/AAAA"),
       ("A5","Taux de change : 1 USD = ... FCFA","B5",600)]
for ac, at, bc, bv in hdr:
    ws[ac] = at; ws[ac].font = BOLD
    ws[bc] = bv; ws[bc].font = BLUE; ws[bc].fill = INPUT_FILL; ws[bc].border = BOX
ws['C5'] = "<- mets le taux du jour"; ws['C5'].font = SMALL

cols = ["Poste de coût","Maersk","CMA CGM","MSC","Intermédiaire 1","Intermédiaire 2"]
for i, h in enumerate(cols, start=1):
    c = ws.cell(7, i, h); c.font = H2
    c.fill = PatternFill("solid", fgColor=VIOLET)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[7].height = 24

def band(r, label, color=GREY_BG, fcol="333333"):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = ws.cell(r, 1, label); c.font = Font(name="Arial", size=10, bold=True, color=fcol)
    c.fill = PatternFill("solid", fgColor=color)

def line(r, label, note=""):
    ws.cell(r, 1, label).font = BLACK; ws.cell(r, 1).border = BOX
    for col in range(2, 7):
        c = ws.cell(r, col); c.font = BLUE; c.fill = INPUT_FILL
        c.border = BOX; c.number_format = '#,##0.00'

band(8, "  A · FRET MARITIME ET SURCHARGES  (en USD)", VIOLET_BG, VIOLET)
postes_a = ["Fret maritime de base","THC au départ (manutention Nansha)","BAF — surcharge carburant",
            "ISPS — surcharge sûreté","CAF — ajustement monétaire","PSS — surcharge haute saison",
            "Frais de documentation / connaissement","Plomb / scellé",
            "Surcharge marchandises dangereuses (UN3481)","THC à l'arrivée (manutention Douala)",
            "Autres surcharges"]
r = 9
for p in postes_a: line(r, p); r += 1
ws.cell(r, 1, "SOUS-TOTAL A").font = BOLD
ws.cell(r, 1).fill = PatternFill("solid", fgColor=GREY_BG); ws.cell(r,1).border = BOX
for col in range(2, 7):
    L = get_column_letter(col)
    c = ws.cell(r, col, f"=SUM({L}9:{L}{r-1})")
    c.font = BOLD; c.fill = PatternFill("solid", fgColor=GREY_BG)
    c.border = BOX; c.number_format = '#,##0.00'
ST_A = r; r += 2

band(r, "  B · OPÉRATIONS EN CHINE  (en USD)", VIOLET_BG, VIOLET); r += 1
start_b = r
for p in ["Camion : entrepôt Guangzhou -> terminal","Déclaration d'exportation chinoise",
          "Manutention / empotage (si sous-traité)","Autres frais au départ"]:
    line(r, p); r += 1
ws.cell(r, 1, "SOUS-TOTAL B").font = BOLD
ws.cell(r, 1).fill = PatternFill("solid", fgColor=GREY_BG); ws.cell(r,1).border = BOX
for col in range(2, 7):
    L = get_column_letter(col)
    c = ws.cell(r, col, f"=SUM({L}{start_b}:{L}{r-1})")
    c.font = BOLD; c.fill = PatternFill("solid", fgColor=GREY_BG)
    c.border = BOX; c.number_format = '#,##0.00'
ST_B = r; r += 2

ws.cell(r, 1, "TOTAL EN USD").font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
ws.cell(r, 1).fill = PatternFill("solid", fgColor=AMBER)
for col in range(2, 7):
    L = get_column_letter(col)
    c = ws.cell(r, col, f"={L}{ST_A}+{L}{ST_B}")
    c.font = Font(name="Arial", size=11, bold=True); c.fill = PatternFill("solid", fgColor=AMBER_BG)
    c.border = BOX; c.number_format = '#,##0.00'
TOT = r; r += 1
ws.cell(r, 1, "TOTAL EN FCFA").font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
ws.cell(r, 1).fill = PatternFill("solid", fgColor=ORANGE)
for col in range(2, 7):
    L = get_column_letter(col)
    c = ws.cell(r, col, f"={L}{TOT}*$B$5")
    c.font = Font(name="Arial", size=11, bold=True); c.fill = PatternFill("solid", fgColor="FFEDE6")
    c.border = BOX; c.number_format = '#,##0'
r += 2

band(r, "  C · CE QUI N'EST PAS UN PRIX — mais qui décide autant", AMBER_BG, "8A5A00"); r += 1
for q in ["Service direct ? (oui / non)","Nombre de transbordements","Lieu(x) de transbordement",
          "Durée de transit annoncée (jours)","Fréquence des départs (par semaine)",
          "Date de cut-off","Garantie de chargement ? (oui / non)","Prix valable jusqu'au",
          "Jours de franchise au départ","Jours de franchise à l'arrivée (Douala)"]:
    ws.cell(r, 1, q).font = BLACK; ws.cell(r, 1).border = BOX
    for col in range(2, 7):
        c = ws.cell(r, col); c.font = BLUE; c.fill = INPUT_FILL; c.border = BOX
        c.alignment = Alignment(horizontal="center")
    r += 1
r += 1
band(r, "  D · DÉCISION", VIOLET_BG, VIOLET); r += 1
ws.cell(r, 1, "Offre retenue et pourquoi").font = BOLD; ws.cell(r, 1).border = BOX
ws.merge_cells(start_row=r, start_column=2, end_row=r+2, end_column=6)
ws.cell(r, 2).fill = INPUT_FILL; ws.cell(r, 2).font = BLUE
ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")

# ═══════════ ONGLET 4 — HISTORIQUE ═══════════
ws = wb.create_sheet("Historique")
ws['A1'] = "HISTORIQUE DES COTATIONS"; ws['A1'].font = H1
ws['A2'] = "Une ligne par cotation obtenue. Au bout de 30 lignes, tu connais le marché mieux que ton transitaire."
ws['A2'].font = SMALL
heads = ["Date","Réf. dossier","Port départ","Port arrivée","Fournisseur","Type conteneur",
         "Total USD","Transit (jours)","Transbordements","Direct ?","Retenue ?","Notes"]
widths = [12,16,18,16,18,15,12,14,16,10,11,45]
for i,(h,w) in enumerate(zip(heads,widths), start=1):
    c = ws.cell(4, i, h); c.font = H2
    c.fill = PatternFill("solid", fgColor=VIOLET)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[4].height = 26
example = ["07/09/2026","BZ-CT-2604-01","Nansha","Douala","Maersk Spot","40HC",
           3150,38,1,"Non","Oui","EXEMPLE — remplace cette ligne. Transbordement Tanger Med."]
for i, v in enumerate(example, start=1):
    c = ws.cell(5, i, v); c.font = Font(name="Arial", size=10, italic=True, color="8A8A96")
    c.border = BOX
    if i == 7: c.number_format = '#,##0'
for rr in range(6, 40):
    for i in range(1, 13):
        c = ws.cell(rr, i); c.font = BLUE; c.border = BOX
        if i == 7: c.number_format = '#,##0'
ws.cell(41, 6, "Moyenne USD").font = BOLD
c = ws.cell(41, 7, "=IFERROR(AVERAGE(G6:G39),\"\")"); c.font = BOLD; c.number_format = '#,##0'
ws.cell(42, 6, "Prix le plus bas").font = BOLD
c = ws.cell(42, 7, "=IFERROR(MIN(G6:G39),\"\")"); c.font = BOLD; c.number_format = '#,##0'
ws.cell(43, 6, "Prix le plus haut").font = BOLD
c = ws.cell(43, 7, "=IFERROR(MAX(G6:G39),\"\")"); c.font = BOLD; c.number_format = '#,##0'

wb.save(OUT)
print("OK:", OUT)
